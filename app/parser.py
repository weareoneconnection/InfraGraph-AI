"""Structure-aware extraction for engineering document uploads.

The parser keeps plain text for knowledge extraction and page/row/paragraph
blocks so every result can point to a human-readable source location. Empty
PDF pages are marked as requiring OCR; the service never invents OCR output.
"""
from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from io import BytesIO, StringIO
from pathlib import Path

SUPPORTED = ".txt .md .csv .json .pdf .docx"


@dataclass
class SourceBlock:
    page: int
    index: int
    start: int
    end: int
    text: str
    kind: str = "paragraph"
    locator: str = ""

    def dump(self) -> dict:
        return asdict(self)


@dataclass
class ParsedDocument:
    text: str
    format: str
    pages: int
    blocks: list[SourceBlock]
    tables: int = 0
    ocr_status: str = "NOT_REQUIRED"
    warnings: list[str] | None = None

    def profile(self) -> dict:
        return {"format": self.format, "pages": self.pages, "blocks": len(self.blocks), "tables": self.tables,
                "ocr_status": self.ocr_status, "warnings": self.warnings or []}


def _assemble(groups: list[tuple[int, str, list[tuple[str, str]]]], fmt: str, warnings=None) -> ParsedDocument:
    parts: list[str] = []
    blocks: list[SourceBlock] = []
    tables = 0
    position = 0
    for page, page_label, items in groups:
        clean = [(kind, " ".join(text.split())) for kind, text in items if text and text.strip()]
        for index, (kind, text) in enumerate(clean, 1):
            if parts:
                parts.append("\n")
                position += 1
            start = position
            parts.append(text)
            position += len(text)
            locator = f"{page_label} · {('表格' if kind == 'table' else '段落')} {index}"
            blocks.append(SourceBlock(page, index, start, position, text, kind, locator))
            tables += kind == "table"
    return ParsedDocument("".join(parts), fmt, max((g[0] for g in groups), default=1), blocks, tables, warnings=warnings or [])


def _text_document(text: str, fmt: str) -> ParsedDocument:
    return _assemble([(1, "正文", [("paragraph", line) for line in text.splitlines() if line.strip()])], fmt)


def parse_document(name: str, content: bytes) -> ParsedDocument:
    ext = Path(name).suffix.lower()
    if ext in {".txt", ".md"}:
        return _text_document(content.decode("utf-8", errors="replace"), ext[1:].upper())
    if ext == ".csv":
        rows = list(csv.reader(StringIO(content.decode("utf-8-sig", errors="replace"))))
        items = [("table", " | ".join(cell.strip() for cell in row)) for row in rows if any(cell.strip() for cell in row)]
        return _assemble([(1, "CSV", items)], "CSV")
    if ext == ".json":
        raw = content.decode("utf-8-sig", errors="replace")
        try:
            raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass
        return _text_document(raw, "JSON")
    if ext == ".pdf":
        from pypdf import PdfReader
        groups, empty = [], []
        for page_no, page in enumerate(PdfReader(BytesIO(content)).pages, 1):
            text = page.extract_text() or ""
            if len(text.strip()) < 10:
                empty.append(page_no)
            groups.append((page_no, f"第 {page_no} 页", [("paragraph", line) for line in text.splitlines()]))
        parsed = _assemble(groups, "PDF", [f"第 {p} 页无可提取文字，需要 OCR" for p in empty])
        parsed.ocr_status = "REQUIRED" if empty else "TEXT_LAYER"
        return parsed
    if ext == ".docx":
        from docx import Document
        doc = Document(BytesIO(content))
        items: list[tuple[str, str]] = [("paragraph", p.text) for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                items.append(("table", " | ".join(cell.text.strip() for cell in row.cells)))
        return _assemble([(1, "Word", items)], "DOCX")
    raise ValueError(f"supported formats: {SUPPORTED}")


def parse_bytes(name: str, content: bytes) -> str:
    """Backwards-compatible plain-text parser used by existing callers."""
    return parse_document(name, content).text
